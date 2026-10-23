"""Verify actual exported files, search, anchors, and mobile layout."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import os
import re
import subprocess
import zipfile
from docx import Document
from pypdf import PdfReader
from playwright.sync_api import sync_playwright
from book_config import CHAPTER_COUNT, STEM

ROOT=Path(__file__).resolve().parents[1]
DIST=ROOT/'dist'


def main():
    report={'run_at_utc':datetime.now(timezone.utc).isoformat(),'status':'passed','checks':[],'files':[]}
    os.environ.setdefault('XDG_CACHE_HOME','/workspace/handbook-tools/cache')
    os.environ.setdefault('XDG_CONFIG_HOME','/workspace/handbook-tools/config')
    html=(DIST/f'{STEM}.html').read_text()
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True)
        page=browser.new_page(viewport={'width':390,'height':844})
        errors=[]
        page.on('pageerror',lambda e:errors.append(str(e)))
        page.set_content(html,wait_until='networkidle')
        page.evaluate('document.fonts.ready')
        audit=page.evaluate('''() => {
          const ids=[...document.querySelectorAll('[id]')].map(x=>x.id);
          const broken=[...document.querySelectorAll('a[href^="#"]')]
            .map(a=>a.getAttribute('href').slice(1)).filter(id=>!document.getElementById(decodeURIComponent(id)));
          return {duplicate:ids.filter((id,i)=>ids.indexOf(id)!==i),broken,
            width:document.documentElement.clientWidth,scroll:document.documentElement.scrollWidth,
            language:document.documentElement.lang};
        }''')
        assert not audit['duplicate'],audit['duplicate']
        assert not audit['broken'],audit['broken']
        assert audit['scroll'] <= audit['width']+1,audit
        assert audit['language']=='zh-TW',audit
        for term in ['管線','Get-ChildItem','物件','Pester','JSON','Compress-Archive','雜湊']:
            page.locator('#book-search').fill(term)
            assert page.locator('#search-results a').count()>0,term
        page.locator('#book-search').fill('不存在xyz123')
        assert page.locator('#search-results a').count()==0
        page.locator('#book-search').fill('Get-ChildItem')
        page.screenshot(path=str(DIST/'驗證_手機搜尋.png'))
        page.locator('#get-childitem').scroll_into_view_if_needed()
        page.screenshot(path=str(DIST/'驗證_手機命令教學.png'))
        page.set_viewport_size({'width':1280,'height':900})
        assert page.evaluate('document.documentElement.scrollWidth <= document.documentElement.clientWidth + 1')
        assert not errors,errors
        browser.close()
    report['checks'].append({'name':'HTML anchors/search/390px+1280px/no-JS-errors','status':'passed','audit':audit})
    docx=DIST/f'{STEM}.docx'
    with zipfile.ZipFile(docx) as z:
        assert z.testzip() is None
        xml=z.read('word/document.xml').decode()
        assert 'Noto Sans CJK TC' in z.read('word/styles.xml').decode()
        assert '<w:hyperlink' in xml
    document=Document(docx)
    text='\n'.join(p.text for p in document.paragraphs)
    for n in range(1,CHAPTER_COUNT+1):assert f'第 {n} 章' in text,n
    assert len(document.tables)>=20
    report['checks'].append({'name':f'DOCX ZIP/XML/reopen/{CHAPTER_COUNT} chapters/tables/fonts/hyperlinks','status':'passed','tables':len(document.tables),'paragraphs':len(document.paragraphs)})
    for filename in [f'{STEM}.pdf',f'{STEM}_手機版.pdf','PowerShell_常用指令速查表_2026.pdf']:
        path=DIST/filename; pdf=PdfReader(path)
        assert len(pdf.pages)>0
        contents=[page.extract_text() for page in pdf.pages]
        assert all(x.strip() for x in contents),'Blank PDF page: '+filename
        joined='\n'.join(contents)
        expected_term='速查' if '速查表' in filename else '管線'
        assert 'PowerShell' in joined and expected_term in joined,filename
        assert not re.search('[\u2e80-\u2fff]',joined), 'Incorrect radical Unicode mapping: '+filename
        if filename.startswith(STEM):
            for n in range(1,CHAPTER_COUNT+1):assert f'第 {n} 章' in joined,(filename,n)
        assert pdf.outline,filename
        fonts=subprocess.run(['pdffonts',str(path)],capture_output=True,text=True,check=True)
        assert not fonts.stderr,fonts.stderr
        rows=fonts.stdout.splitlines()[2:]
        flags=[re.search(r'\b(yes|no)\s+(yes|no)\s+(yes|no)\s+\d+\s+\d+\s*$',row) for row in rows]
        assert rows and all(m and m[1]=='yes' and m[3]=='yes' for m in flags),filename
        width=float(pdf.pages[0].mediabox.width)
        if '手機版' in filename:assert 290<width<305,width
        else:assert 590<width<600,width
        report['checks'].append({'name':filename,'status':'passed','pages':len(pdf.pages),'width_pt':width,'extracted_characters':len(joined),'font_entries':len(rows)})
    for path in sorted(DIST.iterdir()):
        if path.suffix in ('.md','.html','.docx','.pdf') and (path.name.startswith(STEM) or path.name.startswith(('PowerShell_常用','PowerShell_參考'))):
            assert path.stat().st_size>0
            report['files'].append({'name':path.name,'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    (ROOT/'progress/artifact-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'status':'passed','checks':report['checks']},ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
