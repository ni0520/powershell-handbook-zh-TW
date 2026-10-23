"""Build the current authored scope from shared Markdown; no remote services."""
from pathlib import Path
import argparse
import json
import os
import re
import subprocess
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / 'dist'


def run(*args):
    command=list(args)
    if command[0]=='pandoc':
        command += ['--data-dir','assets/pandoc-data','--metadata','toc-title=目錄']
    subprocess.run(command, cwd=ROOT, check=True)
    if command[0]=='pandoc' and '-o' in command:
        output=Path(command[command.index('-o')+1])
        if output.suffix=='.html':
            license_text=(ROOT/'assets/fonts/OFL.txt').read_text()
            output.write_text(output.read_text().replace('</body>',
                '<script type="text/plain" id="font-license">'+license_text+'</script>\n</body>'))


def reference_docx():
    path = DIST / 'reference.docx'
    doc = Document()
    for style in doc.styles:
        if style.type in (1, 2):
            style.font.name = 'Noto Sans CJK TC'
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn('w:eastAsia'), 'Noto Sans CJK TC')
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    section = doc.sections[0]
    section.header.paragraphs[0].text = 'PowerShell 2026｜前 20 章初稿 v0.1'
    footer = section.footer.paragraphs[0]
    footer.text = '頁 '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.save(path)
    return path


def combine(files):
    targets = {}
    for p in (ROOT/'docs').rglob('*.md'):
        anchor = re.search(r'<a id="([^"]+)"></a>', p.read_text())
        if anchor:
            targets[p.resolve()] = anchor[1]
    targets[ROOT/'progress/VALIDATION.md'] = 'delivery-status'
    targets[ROOT/'progress/STATUS.md'] = 'delivery-status'
    chunks = []
    for p in files:
        text = p.read_text()
        # GitHub-readable source anchors become Pandoc anchors for all export formats.
        text = re.sub(r'<a id="([^"]+)"></a>\s*\n(#{1,6} [^\n]+)',
                      lambda m: m[2]+' {#'+m[1]+'}', text)
        def replace_link(match):
            label, link = match[1], match[2]
            if '://' in link or link.startswith('mailto:'):
                return match[0]
            path, _, anchor = link.partition('#')
            target = (p.parent/path).resolve() if path else p.resolve()
            if anchor:
                return f'[{label}](#{anchor})'
            if target in targets:
                return f'[{label}](#{targets[target]})'
            raise ValueError(f'Unmapped book link: {p}: {link}')
        text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', replace_link, text)
        # Short ASCII destinations avoid oversized PDF name tokens from Chinese headings.
        section_number = 0
        def heading_id(match):
            nonlocal section_number
            section_number += 1
            if '{#' in match[0]:
                return match[0]
            return match[0] + f' {{#{p.stem}-s{section_number}}}'
        text = re.sub(r'^#{1,6} .+$', heading_id, text, flags=re.M)
        chunks.append(text)
    chunks.append('''# 交付範圍與驗證說明 {#delivery-status}

本版是前 20 章初稿，含 400 題與解答、10 個核心命令詳解、40 個術語、22 個參數、55 筆官方來源。第 21–40 章、500+ 指令索引及企業案例尚未完成，並非全書定稿。

本機 7.6.8 / Debian 13 的代表性測試與語法紀錄保存在 progress。語法通過不代表每個程式碼區塊都已實際執行。Windows 5.1、Windows/macOS 桌面、企業模組與跨版本效能未實測；PSScriptAnalyzer、Pester 尚未完成可用性驗證，未以其名義宣稱通過。

GitHub 保存狀態以提交後回報與遠端 commit 核對為準；本文件產生時不預先宣稱已合併。
''')
    return '\n\n'.join(chunks)


def add_search(html_path):
    items = []
    for x in json.loads((ROOT/'data/cmdlet-index.json').read_text()):
        items.append({'title':x['name']+' '+x['zh_name'], 'text':x['description']+' '+x['module'], 'anchor':x['id']})
    for x in json.loads((ROOT/'data/glossary.json').read_text()):
        items.append({'title':x['english']+' '+x['chinese'], 'text':x['description'], 'anchor':x['id']})
    for p in sorted((ROOT/'docs/chapters').glob('*.md')):
        title=re.search(r'^# (.+)$',p.read_text(),re.M)[1]
        items.append({'title':title.split(' {#')[0].lstrip('# '), 'text':title, 'anchor':'chapter-'+p.name[:2]})
    panel='''<section class="book-tools" aria-label="書內索引搜尋"><label for="book-search">搜尋章節、指令與中英文術語</label><input id="book-search" type="search" placeholder="例如：管線、Get-ChildItem、錯誤" autocomplete="off"><p>查全文可使用瀏覽器「在網頁中尋找」。本檔可離線閱讀。</p><p id="search-status" role="status" aria-live="polite"></p><ul id="search-results"></ul></section>'''
    script='''<script>
const bookIndex = INDEX;
const input = document.getElementById('book-search');
input.addEventListener('input', () => {
  const q = input.value.trim().toLocaleLowerCase();
  const results = document.getElementById('search-results');
  results.replaceChildren();
  const matches = q ? bookIndex.filter(x => (x.title+' '+x.text).toLocaleLowerCase().includes(q)) : [];
  document.getElementById('search-status').textContent = q ? `找到 ${matches.length} 筆，最多顯示 20 筆` : '';
  for (const entry of matches.slice(0,20)) {
    const li=document.createElement('li'); const a=document.createElement('a');
    a.href='#'+entry.anchor; a.textContent=entry.title; li.append(a); results.append(li);
  }
});
</script>'''.replace('INDEX',json.dumps(items,ensure_ascii=False).replace('</','<\\/'))
    html=html_path.read_text().replace('<body>', '<body>\n'+panel).replace('</body>',script+'\n</body>')
    html_path.write_text(html)


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--sample',action='store_true')
    args=parser.parse_args()
    DIST.mkdir(exist_ok=True)
    run('python3','scripts/prepare_fonts.py')
    # Chromium/fontconfig state uses writable locations, not the read-only home.
    os.environ.setdefault('XDG_CACHE_HOME','/workspace/handbook-tools/cache')
    os.environ.setdefault('XDG_CONFIG_HOME','/workspace/handbook-tools/config')
    chapters=sorted((ROOT/'docs/chapters').glob('*.md'))
    if args.sample:
        files=[chapters[0],chapters[9]]
        stem='排版驗證樣本'
    else:
        files=[ROOT/'docs/README.md',*chapters]
        files += [ROOT/'docs/appendices'/f'{x}.md' for x in ['cmdlet-tutorials','parameter-reference','glossary','master-index']]
        files += [ROOT/'docs/cheatsheets/quick-reference.md',ROOT/'docs/appendices/sources.md']
        stem='PowerShell_學習手冊_2026_前20章初稿'
    md=DIST/f'{stem}.md'; md.write_text(combine(files))
    html=DIST/f'{stem}.html'
    run('pandoc',str(md),'--standalone','--toc','--toc-depth=2','--embed-resources','--metadata','lang=zh-TW','--metadata','title=PowerShell 完整繁體中文版學習手冊（2026）｜前 20 章初稿 v0.1','--css','assets/book.css','-o',str(html))
    if not args.sample:
        add_search(html)
        run('pandoc',str(md),'--standalone','--toc','--toc-depth=2','--metadata','lang=zh-TW','--reference-doc',str(reference_docx()),'-o',str(DIST/f'{stem}.docx'))
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True)
        page=browser.new_page()
        page.set_content(html.read_text(),wait_until='networkidle')
        page.evaluate('document.fonts.ready')
        page.pdf(path=str(DIST/f'{stem}.pdf'),prefer_css_page_size=True,print_background=True,tagged=True,outline=True)
        page.add_style_tag(content='@page{size:105mm 185mm;margin:9mm 8mm}@media print{body{font-size:11pt;line-height:1.65}h1{font-size:18pt}h2{font-size:14pt}h3{font-size:12pt}pre{font-size:8.8pt}table{font-size:8.8pt}th,td{padding:4px}}')
        page.pdf(path=str(DIST/f'{stem}_手機版.pdf'),prefer_css_page_size=True,print_background=True,tagged=True,outline=True)
        page.set_viewport_size({'width':390,'height':844})
        page.emulate_media(media='screen')
        page.screenshot(path=str(DIST/f'{stem}_手機檢查.png'),full_page=False)
        browser.close()
    if not args.sample:
        quick=DIST/'PowerShell_常用指令速查表_2026.md'
        quick.write_text(combine([ROOT/'docs/cheatsheets/quick-reference.md']))
        # Standalone cheat sheet retains source links rather than broken cross-book anchors.
        quick_text=quick.read_text()
        quick_text=re.sub(r'\[([^\]]+)\]\(#[^)]+\)',r'\1（詳主手冊）',quick_text)
        quick.write_text(quick_text)
        quick_html=DIST/'PowerShell_常用指令速查表_2026.html'
        run('pandoc',str(quick),'--standalone','--metadata','title=PowerShell 常用指令速查｜前 20 章','--metadata','lang=zh-TW','--embed-resources','--css','assets/book.css','-o',str(quick_html))
        with sync_playwright() as p:
            browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True)
            page=browser.new_page();page.set_content(quick_html.read_text(),wait_until='networkidle')
            page.evaluate('document.fonts.ready')
            page.pdf(path=str(DIST/'PowerShell_常用指令速查表_2026.pdf'),prefer_css_page_size=True,print_background=True,tagged=True,outline=True)
            browser.close()
        (DIST/'PowerShell_參考資料與來源_2026.md').write_text((ROOT/'docs/appendices/sources.md').read_text().replace('[返回目錄](../README.md)','來源另整合於主手冊。'))
    print(f'Built {stem}')


if __name__=='__main__':
    main()
